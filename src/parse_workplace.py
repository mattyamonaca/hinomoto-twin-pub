"""Normalize the 2020 census 従業地・通学地集計 workbooks (catalog/workplace/manifest.json) into sources/workplace/.

Outputs (all counts are 15歳以上就業者, both sexes; areas are the model leaves: wards of designated cities, other municipalities):
  od_pairs_tidy.csv.gz          origin x destination (leaf or UNK) employed persons, from table 3. The own pair (o,o)
                                 is 自市区町村で従業 (incl. 自宅) plus 従業市区町村「不詳・外国」, which the census counts
                                 at the residence; UNK is 従業地「不詳」.
  od_origin_summary_tidy.csv.gz  per origin: total, own, home, own_out, other, other_in_city, other_in_pref,
                                 other_pref, unknown_muni, unknown_dest (table 3 category rows)
  industry_commute_tidy.csv.gz   table 8: area x employment status x industry (A..T) with the residence-side commuting
                                 breakdown and the workplace-side counts
  odi_residence_tidy.csv.gz      table 9: origin (large cities/wards only) x destination x industry (evaluation only)
  odi_workplace_tidy.csv.gz      table 10: workplace (large cities/wards only) x origin x industry (evaluation only)
  workplace_areas.csv            leaf codes used (from data/geography.csv)
Nothing is imputed; '-' is 0. Aggregate rows (prefectures, designated-city totals, 特別区部) are dropped after checking
that the leaf destinations sum to the published 他市区町村 total.
"""
from paths import SOURCES,RAW,OUTPUT,CATALOG
import json,sys
import numpy as np
import pandas as pd
import openpyxl

DEST=SOURCES/'workplace';WP=RAW/'workplace'
IND=[chr(ord('A')+i) for i in range(20)]          # JSIC major divisions A..T -> G01..G20 (as in sources/industry)
GCODE={l:f'G{i+1:02}' for i,l in enumerate(IND)}
def num(v):
 if v in ('-',None,'',' ','…','X'):return 0.
 return float(v) if isinstance(v,(int,float)) else float(str(v).replace(',',''))
def code_of(cell):return str(cell).split('_')[0]
def leaves():
 g=pd.read_csv(OUTPUT/'geography.csv',dtype=str);pm=pd.read_csv(OUTPUT/'parent_mapping.csv',dtype=str)
 agg=set(pm.parent_code);return [c for c in g.municipality_code if c not in agg]

def manifest():return json.loads((CATALOG/'workplace/manifest.json').read_text(encoding='utf-8'))
def files(table):return [RAW/f['raw_file'].split('/',1)[1] for f in next(s for s in manifest()['sources'] if s['table']==table)['files']]

def read_od(path,value_cols,origin_col,cat_col,other_col):
 """Generic reader for the OD-style tables (3, 9, 10). Yields (origin_code, category_code, other_code, other_level, values)."""
 ws=openpyxl.load_workbook(path,read_only=True).active
 for r in ws.iter_rows(min_row=11,values_only=True):
  if r[origin_col] is None:continue
  yield code_of(r[origin_col]),code_of(r[cat_col]),code_of(r[other_col]),str(r[other_col-2]),[num(r[c]) for c in value_cols]

def od_table3(leaf):
 L=set(leaf);pairs=[];summ={}
 for path in files('3'):
  print('table 3',path.name,flush=True)
  for o,cat,d,lvl,v in read_od(path,[11],3,5,8):
   if o not in L:continue
   x=v[0];s=summ.setdefault(o,{})
   if d=='00000':
    key={'0':'total','01':'own','011':'home','012':'own_out','02':'other','021':'other_in_city','022':'other_in_pref','023':'other_pref','024':'unknown_muni','03':'unknown_dest'}.get(cat)
    if key:s[key]=x
   elif cat=='02' and lvl in ('0','2','3') and d in L and x>0:pairs.append((o,d,x))
   elif cat=='02' and lvl in ('0','2','3') and d not in L and x>0:s['_dropped']=s.get('_dropped',0.)+x   # destinations outside the leaf set (none expected)
 P=pd.DataFrame(pairs,columns=['origin','dest','employed'])
 rows=[]
 for o,s in summ.items():
  own=s.get('own',0.)+s.get('unknown_muni',0.)      # 不詳・外国 is listed under the residence code in the destination list
  rows.append((o,o,own));rows.append((o,'UNK',s.get('unknown_dest',0.)))
 P=pd.concat([P[P.origin!=P.dest],pd.DataFrame(rows,columns=P.columns)],ignore_index=True)
 P=P[P.employed>0].sort_values(['origin','dest']).reset_index(drop=True)
 S=pd.DataFrame([{'origin':o,**{k:v for k,v in s.items()}} for o,s in summ.items()]).fillna(0.)
 for c in ['total','own','home','own_out','other','other_in_city','other_in_pref','other_pref','unknown_muni','unknown_dest','_dropped']:
  if c not in S:S[c]=0.
 S=S.sort_values('origin').reset_index(drop=True)
 # checks: leaf destinations (+ own listing of 不詳・外国) sum to the published 他市区町村 total
 other=P[(P.origin!=P.dest)&(P.dest!='UNK')].groupby('origin').employed.sum().reindex(S.origin).fillna(0).to_numpy()+S.unknown_muni.to_numpy()
 err=np.abs(other-S.other.to_numpy());print('table 3: origins',len(S),'pairs',len(P),'max |sum(leaf dests)+unknown_muni - other| =',float(err.max()),'dropped',float(S._dropped.sum()))
 tot=P.groupby('origin').employed.sum().reindex(S.origin).fillna(0).to_numpy();err2=np.abs(tot-S.total.to_numpy());print('table 3: max |sum(all dests incl UNK) - total| =',float(err2.max()))
 P.to_csv(DEST/'od_pairs_tidy.csv.gz',index=False);S.drop(columns=['_dropped']).to_csv(DEST/'od_origin_summary_tidy.csv.gz',index=False);return P,S

def industry_table8(leaf):
 L=set(leaf);ws=openpyxl.load_workbook(WP/'t8_total.xlsx',read_only=True).active;rows=[]
 cols=['res_total','not_working','own','home','own_out','other','other_in_city','other_in_pref','other_pref','unknown_muni','unknown_dest','outflow','wp_total','wp_from_other','wp_in_city','wp_in_pref','wp_other_pref','wp_unknown_origin','inflow']
 for r in ws.iter_rows(min_row=11,values_only=True):
  if r[0] is None:continue
  area=code_of(r[2])
  if area not in L or code_of(r[3])!='0':continue
  ind=code_of(r[6])
  if ind not in IND and ind!='0':continue
  rows.append((area,code_of(r[4]),GCODE.get(ind,'G00'),*[num(x) for x in r[7:26]]))
 df=pd.DataFrame(rows,columns=['area','status','industry']+cols).drop(columns=['not_working','outflow','inflow'])
 df.to_csv(DEST/'industry_commute_tidy.csv.gz',index=False);print('table 8 rows',len(df),'areas',df.area.nunique());return df

def odi(table,leaf):
 """Tables 9 (residence side) and 10 (workplace side): large-city rows x other municipality x industry."""
 L=set(leaf);recs=[];summ={};vcols=[9]+list(range(11,30))   # A, then B..T (skip 01_うち農業 at column 10)
 for path in files(table):
  print('table',table,path.name,flush=True)
  for a,cat,b,lvl,v in read_od(path,vcols,2,4,7):
   if a not in L:continue
   s=summ.setdefault(a,{})
   if b=='00000':
    if table=='9':key={'0':'total','01':'own','011':'home','012':'own_out','02':'other','021':'other_in_city','022':'other_in_pref','023':'other_pref','024':'unknown_muni','03':'unknown_dest'}.get(cat)
    else:key={'0':'total','01':'own','011':'home','012':'own_out','02':'other','021':'other_in_city','022':'other_in_pref','023':'other_pref','03':'unknown_muni','04':'unknown_dest'}.get(cat)
    if key:s[key]=np.array(v)
   elif cat=='02' and lvl in ('0','2','3') and b in L and sum(v)>0:recs.append((a,b,*v))
 cols=['a','b']+[GCODE[l] for l in IND];P=pd.DataFrame(recs,columns=cols)
 rows=[]
 for a,s in summ.items():
  own=s.get('own',np.zeros(20))+s.get('unknown_muni',np.zeros(20))
  rows.append((a,a,*own))
  if table=='9':rows.append((a,'UNK',*s.get('unknown_dest',np.zeros(20))))
  else:rows.append((a,'UNK',*s.get('unknown_dest',np.zeros(20))))   # residents of the workplace city whose workplace is unknown, counted here by the census
 P=pd.concat([P[P.a!=P.b],pd.DataFrame(rows,columns=cols)],ignore_index=True)
 P=P[P[cols[2:]].sum(1)>0].sort_values(['a','b']).reset_index(drop=True)
 P=P.rename(columns={'a':'origin' if table=='9' else 'workplace','b':'dest' if table=='9' else 'origin'})
 tot=np.array([s['total'].sum() for s in summ.values()]);got=P.groupby(P.columns[0])[cols[2:]].sum().sum(1).reindex(list(summ)).fillna(0).to_numpy()
 print('table',table,'rows',len(P),'areas',len(summ),'max |sum dests - total| =',float(np.abs(got-tot).max()))
 P.to_csv(DEST/('odi_residence_tidy.csv.gz' if table=='9' else 'odi_workplace_tidy.csv.gz'),index=False);return P

def main(tables=None):
 DEST.mkdir(parents=True,exist_ok=True);leaf=leaves();pd.DataFrame({'area':leaf}).to_csv(DEST/'workplace_areas.csv',index=False)
 tables=tables or ['3','8','9','10']
 if '3' in tables:od_table3(leaf)
 if '8' in tables:industry_table8(leaf)
 if '9' in tables:odi('9',leaf)
 if '10' in tables:odi('10',leaf)
if __name__=='__main__':main(sys.argv[1:] or None)
