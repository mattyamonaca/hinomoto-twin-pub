"""Normalize the industry tables used by the M2 experiment (census 6-3 and ESS regional table 24).

Industry classes are the 20 JSIC major divisions A-T (G01..G20; G20 = 分類不能の産業). Age 75+ is merged as in
the rest of the model. Raw workbooks: raw/industry/census_industry_age.xlsx (statInfId=000032201184) and
raw/industry/ess_industry_income.xlsx (statInfId=000040077604). Nothing is imputed.
"""
from paths import SOURCES,RAW
import hashlib,json
import numpy as np
import pandas as pd
import openpyxl

DEST=SOURCES/'industry'
INDUSTRY=[('G01','A','農業，林業'),('G02','B','漁業'),('G03','C','鉱業，採石業，砂利採取業'),('G04','D','建設業'),('G05','E','製造業'),('G06','F','電気・ガス・熱供給・水道業'),('G07','G','情報通信業'),('G08','H','運輸業，郵便業'),('G09','I','卸売業，小売業'),('G10','J','金融業，保険業'),('G11','K','不動産業，物品賃貸業'),('G12','L','学術研究，専門・技術サービス業'),('G13','M','宿泊業，飲食サービス業'),('G14','N','生活関連サービス業，娯楽業'),('G15','O','教育，学習支援業'),('G16','P','医療，福祉'),('G17','Q','複合サービス事業'),('G18','R','サービス業（他に分類されないもの）'),('G19','S','公務（他に分類されるものを除く）'),('G20','T','分類不能の産業')]
LETTER={l:c for c,l,_ in INDUSTRY}
def number(v):
 if v in ('-',None,'',' '):return 0.
 if isinstance(v,(int,float)):return float(v)
 return float(str(v).replace(',',''))

def census():
 """Municipality x sex x age(13) x industry employed persons; rows keyed by area code with 75+ merged."""
 ws=openpyxl.load_workbook(RAW/'industry/census_industry_age.xlsx',read_only=True).active
 rows={}
 for r in ws.iter_rows(min_row=11,values_only=True):
  ind=r[5];letter=ind.split('_')[0]
  if letter not in LETTER or r[3] not in ('1_男','2_女'):continue
  code,name=r[2].split('_',1);sex=r[3][0]
  v=np.array([number(x) for x in r[7:24]])                 # 01..17 five-year classes (15-19 ... 95+)
  age13=np.concatenate([v[:12],[v[12:].sum()]])
  for ai in range(13):
   key=(code,name,r[0],sex,f'{ai+1:02}');rows.setdefault(key,{})[LETTER[letter]]=age13[ai]
 out=pd.DataFrame([dict(area=k[0],name=k[1],type=k[2],sex=k[3],age=k[4],**{c:v.get(c,0.) for c,_,_ in INDUSTRY}) for k,v in rows.items()])
 out=out.sort_values(['area','sex','age']);out.to_csv(DEST/'census_industry_age_tidy.csv.gz',index=False)
 return out

def ess():
 """ESS regional table 24: area x sex x status x income x industry (persons). Area codes: 2-digit prefecture (+000), national 00000, cities keep the table code."""
 ws=openpyxl.load_workbook(RAW/'industry/ess_industry_income.xlsx',read_only=True).active
 recs=[]
 for r in ws.iter_rows(min_row=10,values_only=True):
  if r[1] is None:continue
  acode,aname=str(r[1]).split('_',1);sex=str(r[3])[0];status=str(r[5]).split('_')[0];income=str(r[7]).split('_')[0]
  area=('00000' if acode=='00' else acode+'000' if len(acode)==2 else 'C'+acode)
  for gi,(c,_,_) in enumerate(INDUSTRY):recs.append((area,aname,sex,status,income,c,number(r[9+gi])))
  recs.append((area,aname,sex,status,income,'G00',number(r[8])))
 out=pd.DataFrame(recs,columns=['area','name','sex','status','income','industry','count'])
 out.to_csv(DEST/'income_industry_tidy.csv.gz',index=False)
 return out

def main():
 DEST.mkdir(parents=True,exist_ok=True)
 c=census();e=ess()
 pd.DataFrame(INDUSTRY,columns=['industry_code','jsic_letter','industry_label']).to_csv(DEST/'industry_bins.csv',index=False)
 nat=c[(c.area=='00000')]
 print('census rows',len(c),'areas',c.area.nunique(),'national employed 15+ (sum of industries, both sexes):',float(nat[[x[0] for x in INDUSTRY]].to_numpy().sum()))
 print('ess rows',len(e),'areas',e.area.nunique(),'national total known income (status 0, G00):',float(e[(e.area=='00000')&(e.sex=='0')&(e.status=='0')&(e.income!='00')&(e.industry=='G00')]['count'].sum()))
 manifest={'retrieved_date':'2026-09-06','publisher':'総務省統計局・e-Stat','sources':[
  {'title':'2020国勢調査 就業状態等基本集計 表6-3 男女，年齢（5歳階級），産業（大分類）別就業者数及び平均年齢（15歳以上就業者）－全国，都道府県，市区町村','url':'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000032201184&fileKind=0','raw_file':'raw/industry/census_industry_age.xlsx','raw_sha256':hashlib.sha256((RAW/'industry/census_industry_age.xlsx').read_bytes()).hexdigest(),'tidy_file':'sources/industry/census_industry_age_tidy.csv.gz'},
  {'title':'2022就業構造基本調査 地域編 表24 男女、従業上の地位・雇用形態・起業の有無、所得（主な仕事からの年間収入・収益）、産業別人口（有業者）','url':'https://www.e-stat.go.jp/stat-search/file-download?statInfId=000040077604&fileKind=0','raw_file':'raw/industry/ess_industry_income.xlsx','raw_sha256':hashlib.sha256((RAW/'industry/ess_industry_income.xlsx').read_bytes()).hexdigest(),'tidy_file':'sources/industry/income_industry_tidy.csv.gz'}],
  'notes':['Industry = JSIC major divisions A-T; G20 is 分類不能の産業.','Table 24 has no age dimension: industry x income shapes are age-independent in M2.','Census 6-3 counts employed persons (paid + unpaid family workers) by residence.']}
 (DEST/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
